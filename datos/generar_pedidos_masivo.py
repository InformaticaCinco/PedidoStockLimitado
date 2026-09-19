#!/usr/bin/env python3
from datetime import datetime, timedelta
from pathlib import Path
import argparse
import random

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

SEED = 20260918
HEADERS = [
    "SOLICITUDID","CLIENTEID","ALMACENID","SKU",
    "CANTIDAD","ZONAENTREGA","FECHASOLICITUD"
]
ZONES = ["LIMA_METROPOLITANA", "LIMA_PROVINCIA", "PROVINCIA"]
CLIENTS = [f"CLI-{i:04d}" for i in range(1, 41)]
VALID_SKUS = [f"SKU-{i:05d}" for i in range(2, 41)]

def warehouse_for_sku(sku: str) -> str:
    n = int(sku.split("-")[1])
    if 1 <= n <= 14:
        return "ALM-01"
    if 15 <= n <= 27:
        return "ALM-02"
    if 28 <= n <= 40:
        return "ALM-03"
    return "ALM-01"

def build_rows():
    rng = random.Random(SEED)
    base_dt = datetime(2026, 9, 18, 9, 0, 0)
    rows = []

    sku_cycle = VALID_SKUS[:]
    rng.shuffle(sku_cycle)
    for i in range(1, 403):
        sku = sku_cycle[(i - 1) % len(sku_cycle)]
        rows.append([
            f"MAS-N-{i:04d}",
            CLIENTS[rng.randrange(len(CLIENTS))],
            warehouse_for_sku(sku),
            sku,
            1,
            ZONES[rng.randrange(len(ZONES))],
            base_dt + timedelta(minutes=i),
        ])

    dup_payloads = []
    for i in range(1, 9):
        sku = f"SKU-{i+1:05d}"
        row = [
            f"MAS-DUP-{i:02d}", f"CLI-{i:04d}", warehouse_for_sku(sku), sku, 1,
            ZONES[(i - 1) % 3], base_dt + timedelta(hours=10, minutes=i)
        ]
        rows.append(row)
        dup_payloads.append(list(row))

    conf_payloads = []
    for i in range(1, 5):
        sku = f"SKU-{i+9:05d}"
        row = [
            f"MAS-CONF-{i:02d}", f"CLI-{i+8:04d}", warehouse_for_sku(sku), sku, 1,
            ZONES[i % 3], base_dt + timedelta(hours=11, minutes=i)
        ]
        rows.append(row)
        conf_payloads.append(list(row))

    for i in range(1, 13):
        rows.append([
            f"MAS-SKUX-{i:02d}", CLIENTS[(i + 12) % 40], "ALM-01",
            f"SKU-{90000+i:05d}", 1, ZONES[(i - 1) % 3],
            base_dt + timedelta(hours=12, minutes=i)
        ])

    for i in range(1, 11):
        sku = VALID_SKUS[(i + 10) % len(VALID_SKUS)]
        rows.append([
            f"MAS-CLIX-{i:02d}", f"CLI-{9000+i:04d}", warehouse_for_sku(sku),
            sku, 1, ZONES[i % 3], base_dt + timedelta(hours=13, minutes=i)
        ])

    rows.extend([
        ["MAS-INV-01","CLI-0001","ALM-01","SKU-00002",0,"LIMA_METROPOLITANA",base_dt+timedelta(hours=14,minutes=1)],
        ["MAS-INV-02","CLI-0002","ALM-01","SKU-00003",-1,"LIMA_METROPOLITANA",base_dt+timedelta(hours=14,minutes=2)],
        ["MAS-INV-03","CLI-0003","ALM-01","SKU-00004",1,"ZONA_DESCONOCIDA",base_dt+timedelta(hours=14,minutes=3)],
        ["MAS-INV-04","","ALM-01","SKU-00005",1,"LIMA_PROVINCIA",base_dt+timedelta(hours=14,minutes=4)],
        ["MAS-INV-05","CLI-0005","ALM-01","",1,"PROVINCIA",base_dt+timedelta(hours=14,minutes=5)],
        ["MAS-INV-06","CLI-0006","","SKU-00006",1,"LIMA_METROPOLITANA",base_dt+timedelta(hours=14,minutes=6)],
        ["","CLI-0007","ALM-01","SKU-00007",1,"LIMA_PROVINCIA",base_dt+timedelta(hours=14,minutes=7)],
        ["MAS-INV-08"," CLI-0008","ALM-01","SKU-00008",1,"PROVINCIA",base_dt+timedelta(hours=14,minutes=8)],
        ["MAS-INV-09","CLI-0009","ALM-01","SKU-00009 ",1,"LIMA_METROPOLITANA",base_dt+timedelta(hours=14,minutes=9)],
        ["MAS-INV-10","CLI-0010","ALM1","SKU-00010",1,"LIMA_PROVINCIA",base_dt+timedelta(hours=14,minutes=10)],
        ["MAS-INV-11","CLI-0011","ALM-1","SKU-00011",1,"PROVINCIA",base_dt+timedelta(hours=14,minutes=11)],
        ["MAS-INV-12","CLI-0012","ALM-01","SKU-00012",0,"LIMA_METROPOLITANA",base_dt+timedelta(hours=14,minutes=12)],
        ["MAS-INV-13","CLI-0013","ALM-01","SKU-00013",-5,"LIMA_PROVINCIA",base_dt+timedelta(hours=14,minutes=13)],
        ["MAS-INV-14","CLI-0014","ALM-01","SKU-00014",1,"LIMA METROPOLITANA",base_dt+timedelta(hours=14,minutes=14)],
        ["MAS-INV-15","CLI-0015","ALM-02","SKU-00015",1,"",base_dt+timedelta(hours=14,minutes=15)],
    ])

    for i in range(1, 7):
        sku = VALID_SKUS[(i + 20) % len(VALID_SKUS)]
        cli = CLIENTS[(i + 20) % 40]
        wh = warehouse_for_sku(sku)
        zone = ZONES[(i + 1) % 3]
        dt = base_dt + timedelta(hours=15, minutes=i)
        sid = f"MAS-DSKU-{i:02d}"
        rows.extend([[sid,cli,wh,sku,1,zone,dt],[sid,cli,wh,sku,2,zone,dt]])

    for i in range(1, 6):
        sku = VALID_SKUS[(i + 27) % len(VALID_SKUS)]
        rows.append([
            f"MAS-DATE-{i:02d}", CLIENTS[(i + 30) % 40], warehouse_for_sku(sku),
            sku, 1, ZONES[(i - 1) % 3], f"2026-09-19T0{i}:15:00"
        ])

    for i in range(1, 21):
        rows.append([
            f"MAS-COMP-{i:02d}", f"CLI-{i:04d}", "ALM-01", "SKU-00001",
            1, "LIMA_METROPOLITANA", base_dt + timedelta(hours=16, minutes=i)
        ])

    rows.extend([list(r) for r in dup_payloads])

    for row in conf_payloads:
        changed = list(row)
        changed[4] = 2
        rows.append(changed)

    assert len(rows) == 500
    return rows

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--output", default="data/pedidos-masivo.xlsx")
    args = ap.parse_args()

    rows = build_rows()

    wb = Workbook()
    ws = wb.active
    ws.title = "PEDIDOS"
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="006AA8")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {"A":18,"B":14,"C":12,"D":14,"E":11,"F":24,"G":22}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for cell in ws["G"][1:]:
        if isinstance(cell.value, datetime):
            cell.number_format = "yyyy-mm-dd hh:mm:ss"

    ws.freeze_panes = "A2"
    tab = Table(displayName="PedidosMasivoTable", ref="A1:G501")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Generado {out} con {len(rows)} filas. Semilla={SEED}")

if __name__ == "__main__":
    main()
