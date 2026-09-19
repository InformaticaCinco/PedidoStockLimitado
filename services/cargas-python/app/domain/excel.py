"""Reglas puras, sin consultas ni normalización silenciosa."""
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from hashlib import sha256
from io import BytesIO
from itertools import groupby
import re
from zipfile import ZipFile, ZIP_DEFLATED
from xml.sax.saxutils import escape

from openpyxl import load_workbook

HEADERS = ('SOLICITUDID', 'CLIENTEID', 'ALMACENID', 'SKU', 'CANTIDAD', 'ZONAENTREGA', 'FECHASOLICITUD')
ZONES = {'LIMA_METROPOLITANA', 'LIMA_PROVINCIA', 'PROVINCIA'}
PATTERNS = dict(SOLICITUDID=r'[A-Za-z0-9-]{1,60}', CLIENTEID=r'CLI-[0-9]{4}',
                ALMACENID=r'ALM-[0-9]{2}', SKU=r'SKU-[0-9]{5}')


class WorkbookError(Exception):
    pass


@dataclass
class Row:
    index: int
    values: dict
    errors: list[str] = field(default_factory=list)
    resolved: dict = field(default_factory=dict)
    date: datetime | None = None


def validate(row):
    for key in HEADERS:
        v = row.values[key]
        if v is None or v == '':
            row.errors.append(f'CAMPO_VACIO:{key}')
        elif isinstance(v, str) and v != v.strip():
            row.errors.append(f'ESPACIOS_NO_PERMITIDOS:{key}')
        elif key in PATTERNS and (not isinstance(v, str) or not re.fullmatch(PATTERNS[key], v)):
            row.errors.append(f'FORMATO_{key}_INVALIDO')
        elif key == 'CANTIDAD' and (type(v) is not int or not 1 <= v <= 50):
            row.errors.append('CANTIDAD_INVALIDA')
        elif key == 'ZONAENTREGA' and (not isinstance(v, str) or v not in ZONES):
            row.errors.append('ZONA_INVALIDA')
        elif key == 'FECHASOLICITUD':
            if isinstance(v, datetime) and v.tzinfo is None:
                row.date = v
            elif isinstance(v, str) and re.fullmatch(r'[0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9]{2}:[0-9]{2}:[0-9]{2}', v):
                try:
                    row.date = datetime.strptime(v, '%Y-%m-%dT%H:%M:%S')
                except ValueError:
                    row.errors.append('FECHA_INVALIDA')
            else:
                row.errors.append('FECHA_INVALIDA')
    return row


def parse(content, max_bytes):
    try:
        # Bound decompression before openpyxl; defusedxml disables entity expansion.
        with ZipFile(BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > max_bytes * 20:
                raise WorkbookError('XLSX_EXPANSION_EXCESIVA')
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) != 1:
                raise WorkbookError('HOJA_UNICA_REQUERIDA')
            sheet = book.worksheets[0]
            sheet.reset_dimensions()  # Do not trust spoofed dimension metadata.
            iterator = sheet.iter_rows()
            header = tuple(c.value for c in next(iterator, ()))
            if len(header) != 7 or set(header) != set(HEADERS):
                raise WorkbookError('ENCABEZADOS_INVALIDOS')
            rows = []
            for index, cells in enumerate(iterator, 2):
                if index > 100001 or len(cells) > 7:
                    raise WorkbookError('DIMENSIONES_INVALIDAS')
                values = {key: cells[i].value if i < len(cells) else None for i, key in enumerate(header)}
                row = Row(index, values)
                for cell in cells:
                    if cell.data_type == 'f':
                        row.errors.append('FORMULA_NO_PERMITIDA')
                rows.append(validate(row))
            if not rows:
                raise WorkbookError('ARCHIVO_SIN_FILAS')
            return rows
        finally:
            book.close()
    except WorkbookError:
        raise
    except Exception as exc:
        raise WorkbookError('XLSX_INVALIDO') from exc


def blocks(rows):
    return [list(group) for _, group in groupby(rows, key=lambda r: r.values['SOLICITUDID'])]


def validate_block(rows):
    errors = []
    if not 1 <= len(rows) <= 20:
        errors.append('CANTIDAD_LINEAS_INVALIDA')
    for key in ('CLIENTEID', 'ALMACENID', 'ZONAENTREGA', 'FECHASOLICITUD'):
        vals = [r.date if key == 'FECHASOLICITUD' else r.values[key] for r in rows]
        if any(v != vals[0] for v in vals):
            errors.append('SOLICITUD_INCONSISTENTE')
            break
    skus = [r.values['SKU'] for r in rows]
    if len(set(skus)) != len(skus):
        errors.append('SKU_REPETIDO_EN_SOLICITUD')
    for row in rows:
        row.errors.extend(errors)


def fingerprint(user, rows):
    first = rows[0].values
    # Java String.length() counts UTF-16 code units.
    canonical = f'{len(user.encode("utf-16-le")) // 2}:{user}|{first["ALMACENID"]}|{first["ZONAENTREGA"]}'
    for row in sorted(rows, key=lambda r: r.values['SKU']):
        canonical += f'|{row.values["SKU"]}:{row.values["CANTIDAD"]}'
    return sha256(canonical.encode('utf-8')).hexdigest()


def money(value):
    return Decimal(value).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)


def report(details):
    """Minimal OOXML, written directly to an in-memory ZIP (openpyxl.save uses temp files)."""
    output = BytesIO()
    namespace = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    with ZipFile(output, 'w', ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml',
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '</Types>')
        archive.writestr('_rels/.rels',
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            '</Relationships>')
        archive.writestr('xl/workbook.xml',
            f'<workbook xmlns="{namespace}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<sheets><sheet name="Resultado" sheetId="1" r:id="rId1"/></sheets></workbook>')
        archive.writestr('xl/_rels/workbook.xml.rels',
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
            '</Relationships>')
        with archive.open('xl/worksheets/sheet1.xml', 'w') as sheet:
            sheet.write(f'<worksheet xmlns="{namespace}"><sheetData>'.encode())

            def write_row(index, values):
                sheet.write(f'<row r="{index}">'.encode())
                for column, value in enumerate(values):
                    address = f'{chr(65 + column)}{index}'  # This fixed contract has 11 columns (A-K).
                    if value is None:
                        xml = f'<c r="{address}"/>'
                    elif type(value) is int:
                        xml = f'<c r="{address}" t="n"><v>{value}</v></c>'
                    else:
                        # inlineStr never executes formulas, even for =,+,-,@ prefixes.
                        xml = f'<c r="{address}" t="inlineStr"><is><t xml:space="preserve">{escape(str(value))}</t></is></c>'
                    sheet.write(xml.encode('utf-8'))
                sheet.write(b'</row>')

            write_row(1, ('INDICEORIGINAL', *HEADERS, 'ESTADO', 'MOTIVO', 'IDPEDIDO'))
            for index, detail in enumerate(sorted(details, key=lambda d: d['IndiceOriginal']), 2):
                write_row(index, [detail['IndiceOriginal'], *(detail['Original'][h] for h in HEADERS),
                                 detail['IdEstado'], detail['Log'], detail['IdPedido']])
            sheet.write(b'</sheetData></worksheet>')
    return output.getvalue()
