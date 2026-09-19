import asyncio
from datetime import timedelta
from io import BytesIO
from unittest.mock import patch

import httpx
import pytest
from openpyxl import load_workbook
from pymongo.errors import DuplicateKeyError

from app.api.http import create_app
from app.application.worker import Worker
from app.domain.excel import parse, validate_block, fingerprint
from app.infrastructure.mongo import LeaseLost, now
from tests.conftest import excel, line

pytestmark = pytest.mark.mongo


async def submit(repo, lines):
    return await repo.receive('entrada.xlsx', excel(lines), 'USR-ADMIN', 'corr-it')


async def execute(repo, lines):
    carga_id = await submit(repo, lines)
    assert await Worker(repo).once()
    carga = await repo.db.Carga.find_one({'_id': carga_id})
    assert carga['IdEstadoProceso'] == 'PROCESADA', await repo.db.CargaTarea.find_one({'IdCarga': carga_id})
    return carga


async def test_indexes_gridfs_and_schema(repo):
    carga = await execute(repo, [line(), line(SKU='SKU-00002')])
    indexes = await repo.db.Pedido.index_information()
    assert indexes['IdSolicitud_1']['unique']
    indexes = await repo.db.CargaDetalle.index_information()
    assert indexes['IdCarga_1_IndiceOriginal_1']['unique']
    files = await repo.db['fs.files'].find({'metadata.cargaId': carga['_id']}).to_list()
    assert {f['metadata']['tipo'] for f in files} == {'ENTRADA', 'REPORTE'}
    assert all(f['metadata']['nombreOriginal'].endswith('.xlsx') for f in files)
    pedido = await repo.db.Pedido.find_one({})
    assert pedido['IdEstado'] == 'RECIBIDO' and 'IdWorker' not in pedido
    # Exact ordinary .NET claim filter inspected in MongoRepositorio.cs.
    eligible = {'$or': [{'IdEstado': 'RECIBIDO', 'IdWorker': None},
                        {'IdEstado': 'REQUIERE_REVISION', 'ReintentoSolicitado': True}]}
    assert await repo.db.Pedido.count_documents(eligible) == 1
    assert pedido['EstadoReserva'] == 'NINGUNA' and pedido['AnulacionSolicitada'] is False
    assert pedido['IdUsuarioComprador'] == 'USR-C' and pedido['UsuarioCreacion'] == 'USR-ADMIN'
    assert str(pedido['Subtotal']) == '40.04'
    assert str(pedido['PesoTotal']) == '5.00'
    assert pedido['Envio'] is None and pedido['Total'] is None
    assert pedido['FechaCreacion'].tzinfo and pedido['FechaModificacion'] is None
    details = await repo.db.PedidoDetalle.find({}).to_list()
    assert len(details) == 2 and all(str(d['PrecioUnitario']) == '10.01' for d in details)
    assert all(str(d['PesoUnitario']) == '1.25' and str(d['Subtotal']) == '20.02' for d in details)
    assert await repo.db.Stock.count_documents({}) == 0
    assert carga['CantAceptadas'] == 2
    report = load_workbook(BytesIO(await repo.download(carga['ArchivoReporteId'])))
    assert report.active.max_row == 3
    report.close()


async def test_atomic_claim_and_expired_fencing(repo):
    carga_id = await submit(repo, [line()])
    a, b = await asyncio.gather(repo.claim('worker-a'), repo.claim('worker-b'))
    old = a or b
    assert (a is None) != (b is None)
    await repo.db.CargaTarea.update_one({'_id': old['_id']}, {'$set': {'LeaseHasta': now() - timedelta(seconds=1)}})
    fresh = await repo.claim('worker-c')
    assert fresh['Intentos'] == 2
    with pytest.raises(LeaseLost):
        await repo.set_total(old, 100)
    await repo.set_total(fresh, 1)
    await repo.heartbeat(fresh)
    assert (await repo.db.Carga.find_one({'_id': carga_id}))['TotalFilas'] == 1


async def test_retry_limit_and_final_crash(repo):
    carga_id = await repo.receive('bad.xlsx', b'corrupt', 'USR-ADMIN', 'corr')
    worker = Worker(repo)
    for attempt in range(1, 4):
        await repo.db.CargaTarea.update_one({'IdCarga': carga_id}, {'$set': {'ProximoIntento': now()}})
        assert await worker.once()
        task = await repo.db.CargaTarea.find_one({'IdCarga': carga_id})
        assert task['Intentos'] == attempt
    assert task['EstadoTarea'] == 'FALLIDA'
    assert (await repo.db.Carga.find_one({'_id': carga_id}))['IdEstadoProceso'] == 'ERROR'
    assert not await worker.once()
    second = await submit(repo, [line()])
    await repo.db.CargaTarea.update_one({'IdCarga': second}, {'$set': dict(EstadoTarea='EN_PROCESO', Intentos=3,
                                                                LeaseHasta=now() - timedelta(seconds=1))})
    await repo.reap_exhausted()
    assert (await repo.db.CargaTarea.find_one({'IdCarga': second}))['EstadoTarea'] == 'FALLIDA'
    assert (await repo.db.Carga.find_one({'_id': second}))['IdEstadoProceso'] == 'ERROR'


async def test_same_excel_repeated_blocks_conflict_and_rejected_id_reusable(repo):
    lines = [line('A'), line('A', SKU='SKU-00002'), line('B'),
             line('A'), line('A', SKU='SKU-00002'), line('C'),
             line('A', CANTIDAD=3), line('A', SKU='SKU-00002'),
             line('D', CANTIDAD=0), line('E'), line('D')]
    carga = await execute(repo, lines)
    details = await repo.db.CargaDetalle.find({'IdCarga': carga['_id']}).sort('IndiceOriginal').to_list()
    assert [d['IdEstado'] for d in details] == ['ACEPTADA']*3 + ['DUPLICADA']*2 + ['ACEPTADA'] + ['RECHAZADA']*3 + ['ACEPTADA']*2
    pedido = await repo.db.Pedido.find_one({'IdSolicitud': 'A'})
    assert await repo.db.Pedido.count_documents({'IdSolicitud': 'A'}) == 1
    assert await repo.db.PedidoDetalle.count_documents({'IdPedido': pedido['_id']}) == 2
    assert all([d['Cantidad'] == 2 async for d in repo.db.PedidoDetalle.find({'IdPedido': pedido['_id']})])
    assert details[6]['Log'] == 'SOLICITUD_ID_CON_CONTENIDO_DISTINTO'
    assert details[8]['IdPedido'] is None and details[10]['IdPedido'] is not None
    assert sum([carga['CantAceptadas'], carga['CantRechazadas'], carga['CantDuplicadas']]) == len(lines)


@pytest.mark.parametrize('kind', ['warehouse', 'product', 'ownership', 'client', 'role', 'inactive', 'zone'])
async def test_business_real_order_and_ownership(repo, kind):
    changes = {}
    expected = ''
    if kind == 'warehouse':
        changes = dict(ALMACENID='ALM-99', SKU='SKU-99999', CLIENTEID='CLI-9999')
        expected = 'ALMACEN_NO_EXISTE_O_INACTIVO'
    elif kind == 'product':
        changes = dict(SKU='SKU-99999', CLIENTEID='CLI-9999')
        expected = 'SKU_NO_EXISTE_O_INACTIVO'
    elif kind == 'ownership':
        changes = dict(ALMACENID='ALM-02')
        expected = 'SKU_NO_PERTENECE_AL_ALMACEN'
    elif kind == 'client':
        changes = dict(CLIENTEID='CLI-9999')
        expected = 'CLIENTE_NO_EXISTE'
    elif kind == 'role':
        await repo.db.Usuario.update_one({'_id': 'USR-C'}, {'$set': {'IdRol': 'ROL-A'}})
        expected = 'CLIENTE_NO_ES_COMPRADOR'
    elif kind == 'inactive':
        await repo.db.Usuario.update_one({'_id': 'USR-C'}, {'$set': {'IdEstado': 'INACTIVO'}})
        expected = 'CLIENTE_USUARIO_INACTIVO'
    else:
        await repo.db.Zona.delete_many({})
        expected = 'ZONA_NO_EXISTE_O_INACTIVA'
    carga = await execute(repo, [line(**changes)])
    detail = await repo.db.CargaDetalle.find_one({'IdCarga': carga['_id']})
    assert detail['Log'] == expected and detail['IdPedido'] is None
    assert await repo.db.Pedido.count_documents({}) == 0


async def test_invalid_row_rejects_whole_block_and_continues(repo):
    carga = await execute(repo, [line('A'), line('A', SKU='SKU-00002', CANTIDAD=0), line('B')])
    details = await repo.db.CargaDetalle.find({'IdCarga': carga['_id']}).sort('IndiceOriginal').to_list()
    assert [d['IdEstado'] for d in details] == ['RECHAZADA', 'RECHAZADA', 'ACEPTADA']
    assert details[0]['Log'] == 'SOLICITUD_INVALIDA' and details[0]['IdPedido'] is None
    assert await repo.db.Pedido.count_documents({}) == 1


async def test_transaction_rollback_on_detail_failure(repo):
    carga_id = await submit(repo, [line()])
    carga = await repo.db.Carga.find_one({'_id': carga_id})
    task = await repo.claim('worker')
    # Real server-side collection validator forces the detail insert to fail after Pedido insert.
    await repo.db.command({'collMod': 'PedidoDetalle', 'validator': {'Cantidad': {'$gt': 100}}, 'validationLevel': 'strict'})
    with pytest.raises(Exception):
        await repo.process_block(task, carga, parse(excel([line()]), 100000))
    assert await repo.db.Pedido.count_documents({}) == 0
    assert await repo.db.PedidoDetalle.count_documents({}) == 0
    assert await repo.db.CargaDetalle.count_documents({}) == 0


async def test_concurrent_same_solicitud_unique_and_conflict(repo):
    ids = [await submit(repo, [line()]) for _ in range(2)]
    tasks = [await repo.claim('a'), await repo.claim('b')]
    await asyncio.gather(*(Worker(repo).process(t) for t in tasks))
    states = [d['IdEstado'] async for d in repo.db.CargaDetalle.find({})]
    assert sorted(states) == ['ACEPTADA', 'DUPLICADA']
    assert await repo.db.Pedido.count_documents({}) == 1
    assert await repo.db.PedidoDetalle.count_documents({}) == 1
    pedido = await repo.db.Pedido.find_one({})
    with pytest.raises(DuplicateKeyError):
        await repo.db.Pedido.insert_one(pedido | {'_id': 'PED-other'})
    carga = await execute(repo, [line(CANTIDAD=3)])
    assert carga['CantRechazadas'] == 1


async def test_unique_rows_redelivery_and_crash_after_block(repo):
    carga_id = await submit(repo, [line('A'), line('B')])
    task = await repo.claim('old')
    carga = await repo.db.Carga.find_one({'_id': carga_id})
    rows = parse(excel([line('A')]), 100000)
    await repo.process_block(task, carga, rows)  # Simulated process death after committed first block.
    await repo.db.CargaTarea.update_one({'_id': task['_id']}, {'$set': {'LeaseHasta': now() - timedelta(seconds=1)}})
    assert await Worker(repo).once()
    done = await repo.db.Carga.find_one({'_id': carga_id})
    detail = await repo.db.CargaDetalle.find_one({'IdCarga': carga_id})
    with pytest.raises(DuplicateKeyError):
        await repo.db.CargaDetalle.insert_one(detail | {'_id': 'CD-other'})
    await repo.db.CargaTarea.update_one({'_id': task['_id']}, {'$set': dict(EstadoTarea='EN_PROCESO', LeaseHasta=now() - timedelta(seconds=1))})
    assert await Worker(repo).once()
    again = await repo.db.Carga.find_one({'_id': carga_id})
    assert again['ArchivoReporteId'] == done['ArchivoReporteId']
    assert again['CantAceptadas'] == 2 and again['CantDuplicadas'] == 0
    for name in ('Pedido', 'PedidoDetalle', 'CargaDetalle'):
        assert await repo.db[name].count_documents({}) == 2
    assert await repo.db['fs.files'].count_documents({'metadata.tipo': 'REPORTE'}) == 1


async def test_crash_after_report_before_final_commit_regenerates(repo):
    carga_id = await submit(repo, [line()])
    task = await repo.claim('old')
    original_finish = repo.finish
    async def fail(*args):
        raise RuntimeError('simulated crash')
    with patch.object(repo, 'finish', fail):
        with pytest.raises(RuntimeError):
            await Worker(repo).process(task)
    assert await repo.db['fs.files'].count_documents({'metadata.tipo': 'REPORTE'}) == 1
    await repo.db.CargaTarea.update_one({'_id': task['_id']}, {'$set': {'LeaseHasta': now() - timedelta(seconds=1)}})
    assert await Worker(repo).once()
    carga = await repo.db.Carga.find_one({'_id': carga_id})
    assert carga['IdEstadoProceso'] == 'PROCESADA' and carga['CantAceptadas'] == 1
    assert await repo.db.Pedido.count_documents({}) == 1
    assert await repo.db['fs.files'].count_documents({'metadata.tipo': 'REPORTE'}) == 1


async def test_500_rows_full_http_gridfs_queue_counts_report(repo, auth):
    security, token, _, _ = auth
    app = create_app(repo.settings, repo, security, run_worker=False)
    headers = {'Authorization': 'Bearer ' + token()}
    rows = [line(f'B-{i:04d}', CANTIDAD=0 if i % 10 == 0 else 2) for i in range(500)]
    async with app.router.lifespan_context(app), httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://test') as client:
        result = await client.post('/cargas', files={'archivo': ('500.xlsx', excel(rows))}, headers=headers)
        assert result.status_code == 202
        carga_id = result.json()['data']['cargaId']
        assert await repo.db.Pedido.count_documents({}) == 0  # HTTP never processes rows.
        assert (await client.get(f'/cargas/{carga_id}/reporte', headers=headers)).status_code == 409
        assert (await client.get('/cargas/CG-missing', headers=headers)).status_code == 404
        assert await Worker(repo).once()
        result = (await client.get(f'/cargas/{carga_id}', headers=headers)).json()['data']
        assert result['estado'] == 'PROCESADA'
        assert result['totales'] == dict(filas=500, aceptadas=450, rechazadas=50, duplicadas=0)
        download = await client.get(result['reporte'], headers=headers)
        assert download.status_code == 200 and 'attachment;' in download.headers['content-disposition']
        book = load_workbook(BytesIO(download.content), read_only=True)
        output = list(book.active.values)
        assert len(output) == 501 and [r[0] for r in output[1:]] == list(range(2, 502))
        book.close()
        assert await repo.db.CargaDetalle.count_documents({'IdCarga': carga_id}) == 500
        assert await repo.db.Pedido.count_documents({}) == 450
        assert await repo.db.Stock.count_documents({}) == 0
        carga = await repo.db.Carga.find_one({'_id': carga_id})
        await repo.files.delete(carga['ArchivoReporteId'])
        assert (await client.get(result['reporte'], headers=headers)).status_code == 500


async def test_lease_fences_order_and_backoff(repo):
    carga_id = await submit(repo, [line()])
    task = await repo.claim('expired')
    carga = await repo.db.Carga.find_one({'_id': carga_id})
    await repo.db.CargaTarea.update_one({'_id': task['_id']}, {'$set': {'LeaseHasta': now() - timedelta(seconds=1)}})
    with pytest.raises(LeaseLost):
        await repo.process_block(task, carga, parse(excel([line()]), 100000))
    assert await repo.db.Pedido.count_documents({}) == 0
    fresh = await repo.claim('new')
    await repo.fail(fresh, 'PRUEBA_SEGURA')
    await repo.db.CargaTarea.update_one({'_id': task['_id']}, {'$set': {'ProximoIntento': now() + timedelta(seconds=20)}})
    assert await repo.claim('too-soon') is None
    assert (await repo.db.CargaTarea.find_one({'_id': task['_id']}))['Intentos'] == 2


async def test_invalid_headers_never_create_partial_orders(repo):
    carga_id = await repo.receive('headers.xlsx', excel([line()], ('SOLICITUDID', 'SKU')), 'USR-ADMIN', 'headers')
    assert await Worker(repo).once()
    assert await repo.db.Pedido.count_documents({}) == 0
    assert await repo.db.CargaDetalle.count_documents({}) == 0
    task = await repo.db.CargaTarea.find_one({'IdCarga': carga_id})
    assert task['EstadoTarea'] == 'PENDIENTE' and task['UltimoError'] == 'ENCABEZADOS_INVALIDOS'


async def test_twenty_competing_orders_do_not_read_or_change_stock(repo):
    await repo.db.Stock.insert_many([dict(_id=f'STK-{i}', IdAlmacen='ALM-01', IdProducto=f'SKU-{i:05d}', CantidadStock=1, ReservaStock=0) for i in range(1, 4)])
    before = await repo.db.Stock.find({}).sort('_id').to_list()
    carga = await execute(repo, [line(f'COMPETE-{i}', SKU=f'SKU-{i % 3 + 1:05d}', CANTIDAD=50) for i in range(20)])
    assert carga['CantAceptadas'] == 20
    assert await repo.db.Stock.find({}).sort('_id').to_list() == before
    assert await repo.db.Pedido.count_documents({'IdEstado': 'RECIBIDO'}) == 20
