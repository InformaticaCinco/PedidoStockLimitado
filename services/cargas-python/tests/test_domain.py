from datetime import datetime
from hashlib import sha256
from io import BytesIO
import pytest
from openpyxl import load_workbook

from app.domain.excel import Row, validate, parse, blocks, validate_block, fingerprint, report, WorkbookError, HEADERS
from tests.conftest import line, excel


@pytest.mark.parametrize('key,value', [
    ('SOLICITUDID', 'bad_1'), ('SOLICITUDID', 'x'*61), ('SOLICITUDID', ''),
    ('CLIENTEID', 'CLI-1'), ('CLIENTEID', 1), ('ALMACENID', 'ALM-1'), ('SKU', 'SKU-1'),
    ('CANTIDAD', 0), ('CANTIDAD', -1), ('CANTIDAD', 1.5), ('CANTIDAD', 2.0),
    ('CANTIDAD', True), ('CANTIDAD', '2'), ('CANTIDAD', 51),
    ('ZONAENTREGA', 'OTRA'), ('FECHASOLICITUD', '18/09/2026'),
    ('FECHASOLICITUD', '2026-09-18 10:00:00'), ('FECHASOLICITUD', '2026-02-30T10:00:00'),
    ('FECHASOLICITUD', '2026-09-18T10:00:00Z'), ('FECHASOLICITUD', None),
    *[(k, ' ' + str(line()[k])) for k in HEADERS]])
def test_invalid_scalar(key, value):
    row = validate(Row(2, line(**{key: value})))
    assert row.errors
    assert row.values[key] == value


@pytest.mark.parametrize('quantity', [1, 50])
@pytest.mark.parametrize('date', ['2026-09-18T10:00:00', datetime(2026, 9, 18, 10)])
def test_valid(quantity, date):
    assert not validate(Row(2, line(CANTIDAD=quantity, FECHASOLICITUD=date))).errors


def test_consecutive_not_global():
    rows = parse(excel([line('A'), line('A', SKU='SKU-00002'), line('B'), line('A')]), 100000)
    assert [len(b) for b in blocks(rows)] == [2, 1, 1]
    assert [r.index for r in rows] == [2, 3, 4, 5]


@pytest.mark.parametrize('changes', [dict(SKU='SKU-00001'), dict(CLIENTEID='CLI-0002', SKU='SKU-00002'),
    dict(ALMACENID='ALM-02', SKU='SKU-00002'), dict(ZONAENTREGA='PROVINCIA', SKU='SKU-00002'),
    dict(FECHASOLICITUD='2026-09-19T10:00:00', SKU='SKU-00002')])
def test_block_consistency(changes):
    rows = parse(excel([line(), line(**changes)]), 100000)
    validate_block(rows)
    assert all(r.errors for r in rows)


def test_fingerprint_java_golden_order_identity_quantity():
    rows = parse(excel([line(SKU='SKU-00002'), line()]), 100000)
    expected = sha256(b'5:USR-C|ALM-01|LIMA_METROPOLITANA|SKU-00001:2|SKU-00002:2').hexdigest()
    assert fingerprint('USR-C', rows) == expected
    assert fingerprint('USR-C', list(reversed(rows))) == expected
    assert fingerprint('OTRO', rows) != expected
    rows[0].values['CANTIDAD'] = 3
    assert fingerprint('USR-C', rows) != expected


def test_report_order_formula_and_original():
    details = [dict(IndiceOriginal=i, Original=line(SOLICITUDID=text), IdEstado='RECHAZADA', Log='FORMULA', IdPedido=None)
               for i, text in [(3, '=1+1'), (2, '+SUM(A1)'), (4, '-1'), (5, '@abc')]]
    book = load_workbook(BytesIO(report(details)), data_only=False)
    assert [r[0].value for r in list(book.active.rows)[1:]] == [2, 3, 4, 5]
    assert all(book.active.cell(i, 2).data_type == 's' for i in range(2, 6))
    assert book.active.cell(3, 2).value == '=1+1'
    book.close()


@pytest.mark.parametrize('headers', [HEADERS[:-1], (*HEADERS[:-1], HEADERS[0]), (*HEADERS, 'PRECIO')])
def test_bad_headers(headers):
    with pytest.raises(WorkbookError):
        parse(excel([line()], headers), 100000)


def test_mapped_headers_and_formula_rejection():
    assert not parse(excel([line()], tuple(reversed(HEADERS))), 100000)[0].errors
    assert 'FORMULA_NO_PERMITIDA' in parse(excel([line(CANTIDAD='=1+1')]), 100000)[0].errors


def test_corrupt_and_empty():
    for content in (b'bad', excel([])):
        with pytest.raises(WorkbookError):
            parse(content, 100000)


def test_line_limit_and_no_hidden_rows_omitted():
    rows = parse(excel([line(SKU=f'SKU-{i:05d}') for i in range(21)]), 100000)
    validate_block(rows)
    assert all('CANTIDAD_LINEAS_INVALIDA' in r.errors for r in rows)
    book = load_workbook(BytesIO(excel([line(), line('B')])))
    book.active.row_dimensions[3].hidden = True
    out = BytesIO()
    book.save(out)
    book.close()
    assert len(parse(out.getvalue(), 100000)) == 2


def test_multiple_sheets_and_decompression_bound():
    book = load_workbook(BytesIO(excel([line()])))
    book.create_sheet('extra')
    out = BytesIO()
    book.save(out)
    book.close()
    with pytest.raises(WorkbookError, match='HOJA_UNICA'):
        parse(out.getvalue(), 100000)
    with pytest.raises(WorkbookError, match='EXPANSION'):
        parse(excel([line()]), 1)


def test_report_never_uses_openpyxl_tempfile(monkeypatch):
    from openpyxl.worksheet import _writer
    def forbidden(*args, **kwargs):
        raise AssertionError('No filesystem allowed for operational reports')
    monkeypatch.setattr(_writer, 'create_temporary_file', forbidden)
    detail = dict(IndiceOriginal=2, Original=line(), IdEstado='ACEPTADA', Log='OK', IdPedido='PED-test')
    book = load_workbook(BytesIO(report([detail])), read_only=True)
    assert list(book.active.values)[1][-1] == 'PED-test'
    book.close()
